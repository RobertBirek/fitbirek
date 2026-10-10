#!/usr/bin/env python3
"""Apply the versioned ID → Polish name mapping to BAZA_GLOWNA only.

Run xlsx_to_json.py afterwards, then this script with --check.
Requires openpyxl; no dependencies are installed by this script.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
import json
from pathlib import Path
import tempfile

import openpyxl

import xlsx_to_json as converter

MAPPING_PATH = Path(__file__).parent / "source_data" / "exercise_names_pl.json"


def require(condition, message):
    if not condition:
        raise ValueError(message)


def unique_pairs(pairs):
    result = {}
    for key, value in pairs:
        require(key not in result, f"Duplicate mapping key: {key}")
        result[key] = value
    return result


def collisions(entries):
    groups = defaultdict(list)
    for entry in entries:
        key = (entry["nazwaPl"], entry["partiaGlowna"], entry["typ"])
        groups[key].append(entry["id"])
    return [ids for ids in groups.values() if len(ids) > 1]


def workbook_snapshot(workbook):
    return {
        sheet.title: {
            cell.coordinate: (cell.value, cell.data_type, cell._style,
                              cell.comment, cell.hyperlink)
            for row in sheet for cell in row
        }
        for sheet in workbook
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Verify without writing")
    args = parser.parse_args()
    names = json.loads(MAPPING_PATH.read_text(encoding="utf-8"),
                       object_pairs_hook=unique_pairs)
    expected_ids = {f"cw{i:03d}" for i in range(1, 317)}
    require(set(names) == expected_ids, "Mapping must contain exactly cw001–cw316")
    require(all(isinstance(n, str) and n.strip() == n and n and
                not n.startswith("=") for n in names.values()), "Invalid target name")
    baseline = json.loads(converter.OUTPUT_PATH.read_text(encoding="utf-8"))
    before = converter.convert()
    require(before == baseline, "Baseline export differs from XLSX; refusing to write")
    require({e["id"] for e in before} == expected_ids, "Unexpected source IDs")
    expected = [dict(entry, nazwaPl=names[entry["id"]]) for entry in before]
    changed = sum(a != b for a, b in zip(before, expected))
    print(f"Records: {len(before)}; changed names: {changed}")
    print(f"PL+partia+typ collisions before: {collisions(before)}")
    print(f"PL+partia+typ collisions after: {collisions(expected)}")
    if args.check:
        require(before == expected, "XLSX/export names differ from mapping")
        print("CHECK OK: mapping = XLSX = exported JSON")
        return
    if not changed:
        print("Already up to date; XLSX not rewritten")
        return

    workbook = openpyxl.load_workbook(converter.XLSX_PATH)
    snapshot = workbook_snapshot(workbook)
    sheet = workbook[converter.SHEET_NAME]
    require([c.value for c in sheet[1]] == converter.EXPECTED_HEADER,
            "Unexpected XLSX header")
    id_column = converter.EXPECTED_HEADER.index("ID") + 1
    name_column = converter.EXPECTED_HEADER.index("Nazwa_PL") + 1
    for row in range(2, sheet.max_row + 1):
        number = sheet.cell(row, id_column).value
        if number is None:
            continue
        cell = sheet.cell(row, name_column)
        name = names[f"cw{number:03d}"]
        cell.value = name
        old = snapshot[sheet.title][cell.coordinate]
        snapshot[sheet.title][cell.coordinate] = (name, *old[1:])

    original_path = converter.XLSX_PATH
    with tempfile.NamedTemporaryFile(dir=original_path.parent, suffix=".xlsx",
                                     delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        saved = openpyxl.load_workbook(temporary_path)
        require(workbook_snapshot(saved) == snapshot,
                "Unexpected cell value, type, style, comment or hyperlink change")
        saved.close()
        converter.XLSX_PATH = temporary_path
        after = converter.convert()
        require(after == expected, "Export changed fields other than nazwaPl")
        require({e["id"] for e in after} == expected_ids, "Export changed IDs")
        require(json.loads(converter.OUTPUT_PATH.read_text(encoding="utf-8")) == baseline,
                "JSON changed concurrently; refusing to replace XLSX")
        temporary_path.chmod(original_path.stat().st_mode & 0o777)
        temporary_path.replace(original_path)
    finally:
        converter.XLSX_PATH = original_path
        workbook.close()
        temporary_path.unlink(missing_ok=True)
    print("UPDATED: only Nazwa_PL cells; all other cells and exported fields verified")
    print("Next: run tools/xlsx_to_json.py, then this script with --check")


if __name__ == "__main__":
    main()
