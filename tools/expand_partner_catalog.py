#!/usr/bin/env python3
"""Append approved partner/general movements to XLSX and explicit static inputs.

Does not edit generated assets. Follow with xlsx_to_json.py and metadata builder.
"""
import argparse
from copy import copy
import json
import tempfile
from pathlib import Path

import openpyxl
import xlsx_to_json as converter
from build_catalog_metadata import MUSCLES
from build_execution_conditions import INPUT_PATH, validate_classification
from expand_catalog import SOURCE_PATH, ledger
from partner_additions import ADDITIONS, EXPECTED_COUNT, SOURCES
from update_exercise_names_pl import MAPPING_PATH, workbook_snapshot


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    before = converter.convert()
    if len(before) not in (369, EXPECTED_COUNT):
        raise ValueError('Expected unchanged baseline 369 or completed expansion')
    labels = {i: label for i, label, _ in MUSCLES}
    names = json.loads(MAPPING_PATH.read_text())
    execution = json.loads(INPUT_PATH.read_text())
    validate_classification(execution, {e['id'] for e in before})
    wb = openpyxl.load_workbook(converter.XLSX_PATH)
    original = workbook_snapshot(wb)
    ws = wb[converter.SHEET_NAME]
    existing = {r[0]: list(r) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    added = 0
    for a in ADDITIONS:
        row = [a['number'], a['pl'], a['en'], ', '.join(labels[m] for m in a['primary']),
               ', '.join(labels[m] for m in a['secondary']) or '-', a['equipment'],
               a['kind'], a['level'], a['pattern'], a['prescription'], a['tempo'],
               a['cues'], a['errors'], a['progression'], a['regression'], SOURCES[a['source']]]
        eid = f'cw{a["number"]}'
        if a['number'] in existing:
            if existing[a['number']] != row:
                raise ValueError(f'Existing row differs: {eid}')
        elif args.check:
            raise ValueError(f'Missing row: {eid}')
        else:
            ws.append(row)
            added += 1
            for cell, template in zip(ws[ws.max_row], ws[2]):
                cell._style = copy(template._style)
        for mapping, value in [(names, a['pl']), (execution['exercises'],
                                               {k: a[k] for k in ('conditions', 'positions')})]:
            if eid in mapping and mapping[eid] != value:
                raise ValueError(f'Existing input differs: {eid}')
            if args.check and eid not in mapping:
                raise ValueError(f'Missing input: {eid}')
            mapping[eid] = value
    validate_classification(execution, {f'cw{n:03d}' for n in range(1, EXPECTED_COUNT + 1)})
    texts = {p: json.dumps(v, ensure_ascii=False, indent=2) + '\n'
             for p, v in [(MAPPING_PATH, names), (INPUT_PATH, execution), (SOURCE_PATH, ledger())]}
    if args.check:
        for path, text in texts.items():
            if path.read_text() != text:
                raise ValueError(f'Static input differs: {path}')
    else:
        if added:
            with tempfile.NamedTemporaryFile(dir=converter.XLSX_PATH.parent, suffix='.xlsx', delete=False) as f:
                temp = Path(f.name)
            try:
                wb.save(temp)
                saved = openpyxl.load_workbook(temp)
                snapshot = workbook_snapshot(saved)
                saved.close()
                for sheet, cells in original.items():
                    for coord, value in cells.items():
                        if snapshot[sheet].get(coord) != value:
                            raise ValueError(f'Existing cell changed: {sheet}!{coord}')
                with_converter_path = converter.XLSX_PATH
                try:
                    converter.XLSX_PATH = temp
                    after = converter.convert()
                finally:
                    converter.XLSX_PATH = with_converter_path
                if after[:len(before)] != before or len(after) != EXPECTED_COUNT:
                    raise ValueError('Existing records changed')
                temp.chmod(with_converter_path.stat().st_mode & 0o777)
                temp.replace(with_converter_path)
            finally:
                temp.unlink(missing_ok=True)
        for path, text in texts.items():
            if path.read_text() != text:
                path.write_text(text)
    wb.close()
    print(f'Partner expansion OK: {added} rows appended; {EXPECTED_COUNT} total')


if __name__ == '__main__':
    main()
