"""Catalog data contracts; Python 3.12 + requirements-catalog.lock, no network."""
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import xlsx_to_json as converter

ROOT = Path(__file__).resolve().parents[2]


class ConverterTests(unittest.TestCase):
    def test_new_equipment(self):
        self.assertEqual(converter.normalize_sprzet('Sandbag, Kółko do brzucha'),
                         ['Sandbag', 'Kółko do brzucha'])

    def test_append_without_fixed_count(self):
        import openpyxl
        wb = openpyxl.load_workbook(converter.XLSX_PATH)
        sheet = wb[converter.SHEET_NAME]
        row = [c.value for c in sheet[2]]
        row[0] = sheet.max_row
        row[1:3] = ['Test nowego ruchu', 'New movement test']
        sheet.append(row)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'catalog.xlsx'
            wb.save(path)
            with patch.object(converter, 'XLSX_PATH', path):
                data = converter.convert()
            self.assertEqual(len(data), sheet.max_row - 1)
        wb.close()

    def test_missing_legacy_id_rejected(self):
        self._bad_workbook(lambda s: setattr(s['A2'], 'value', 999), 'ID')

    def test_duplicate_id_rejected(self):
        self._bad_workbook(lambda s: setattr(s['A3'], 'value', 1), 'ID')

    def test_missing_required_field_rejected(self):
        self._bad_workbook(lambda s: setattr(s['P2'], 'value', None), 'Zrodlo')

    def test_legacy_content_change_rejected(self):
        self._bad_workbook(lambda s: setattr(s['B2'], 'value', 'Zmieniona stara nazwa'), 'cw001')

    def test_formula_rejected(self):
        self._bad_workbook(lambda s: setattr(s['B2'], 'value', '=1+1'), 'Nazwa_PL')

    def test_empty_id_with_content_rejected(self):
        self._bad_workbook(lambda s: setattr(s['A2'], 'value', None), 'ID')

    def test_id_gap_rejected(self):
        self._bad_workbook(lambda s: setattr(s.cell(s.max_row, 1), 'value', 998), 'ID')

    def _bad_workbook(self, mutate, message):
        import openpyxl
        wb = openpyxl.load_workbook(converter.XLSX_PATH)
        mutate(wb[converter.SHEET_NAME])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'bad.xlsx'
            wb.save(path)
            with patch.object(converter, 'XLSX_PATH', path):
                with self.assertRaisesRegex(ValueError, message):
                    converter.convert()
        wb.close()


class CatalogTests(unittest.TestCase):
    def setUp(self):
        import validate_catalog
        self.validator = validate_catalog
        self.data = validate_catalog.load(converter.OUTPUT_PATH)
        self.meta = validate_catalog.load(validate_catalog.METADATA_PATH)

    def test_legacy_records_unchanged(self):
        legacy = json.dumps(self.data[:316], ensure_ascii=False, sort_keys=True).encode()
        self.assertEqual(hashlib.sha256(legacy).hexdigest(), self.validator.LEGACY_SHA256)

    def test_export_reproducible(self):
        exported = json.loads(converter.OUTPUT_PATH.read_text())
        self.assertEqual(converter.convert(), exported)

    def test_metadata_and_sources(self):
        import validate_catalog
        validate_catalog.validate()

    def test_metadata_rejects_unknown_muscle(self):
        import validate_catalog
        data = json.loads(converter.OUTPUT_PATH.read_text())
        meta = json.loads(validate_catalog.METADATA_PATH.read_text())
        bad = copy.deepcopy(meta)
        bad['exercises']['cw001']['secondaryMuscles'] = ['cardio']
        with self.assertRaisesRegex(ValueError, 'muscle'):
            validate_catalog.validate_metadata(data, bad)

    def test_names_mapping_check(self):
        subprocess.run([sys.executable, str(ROOT / 'tools/update_exercise_names_pl.py'),
                        '--check'], check=True, capture_output=True)

    def test_metadata_reproducible(self):
        from build_catalog_metadata import build
        self.assertEqual(build(), self.meta)

    def test_reviewed_ambiguous_anatomy(self):
        entries = self.meta['exercises']
        self.assertEqual(entries['cw032']['primaryMuscles'], ['biceps', 'back'])
        self.assertEqual(entries['cw063']['primaryMuscles'], ['hamstrings', 'glutes'])
        self.assertEqual(entries['cw099']['primaryMuscles'], ['brachialis'])
        self.assertEqual(entries['cw168']['primaryMuscles'], ['tibialis_anterior'])
        self.assertEqual(entries['cw300']['primaryMuscles'], ['diaphragm'])
        self.assertEqual(entries['cw357']['styles'], ['calisthenics', 'yoga'])

    def test_direction_is_not_body_side(self):
        for eid in ('cw269', 'cw270', 'cw273', 'cw274', 'cw329'):
            self.assertFalse(self.meta['exercises'][eid]['unilateral'], eid)

    def test_bilateral_sissy_squat_is_not_unilateral(self):
        self.assertFalse(self.meta['exercises']['cw140']['unilateral'])

    def test_metadata_rejects_invalid_contracts(self):
        for key, value in [('tracking', 'distance'), ('unilateral', 1),
                           ('styles', ['unknown']), ('primaryMuscles', []),
                           ('aliases', ['duplicate', 'duplicate']),
                           ('secondaryMuscles', ['chest'])]:
            with self.subTest(key=key):
                bad = copy.deepcopy(self.meta)
                bad['exercises']['cw001'][key] = value
                with self.assertRaises(ValueError):
                    self.validator.validate_metadata(self.data, bad)
        bad = copy.deepcopy(self.meta)
        del bad['exercises']['cw316']
        with self.assertRaisesRegex(ValueError, 'coverage'):
            self.validator.validate_metadata(self.data, bad)
        bad = copy.deepcopy(self.meta)
        bad['schemaVersion'] = True
        with self.assertRaisesRegex(ValueError, 'schemaVersion'):
            self.validator.validate_metadata(self.data, bad)

    def test_source_mismatch_rejected(self):
        sources = self.validator.load(self.validator.SOURCES_PATH)
        bad = copy.deepcopy(self.data)
        bad[316]['zrodlo'] = 'https://example.org/unverified'
        with self.assertRaisesRegex(ValueError, 'source differs'):
            self.validator.validate_sources(bad, sources)

    def test_expansion_input_matches_workbook(self):
        subprocess.run([sys.executable, str(ROOT / 'tools/expand_catalog.py'),
                        '--check'], check=True, capture_output=True)

    def test_equipment_constants_match_converter(self):
        import re
        text = (ROOT / 'lib/app/constants.dart').read_text()
        block = text.split('dostepnySprzetOpcje = [', 1)[1].split('];', 1)[0]
        self.assertEqual(set(re.findall("'([^']+)'", block)), converter.VALID_SPRZET)

    def test_new_ids_groups_and_tracking(self):
        from catalog_additions import ADDITIONS
        self.assertEqual([a['number'] for a in ADDITIONS], list(range(317, 364)))
        from collections import Counter
        self.assertEqual(Counter(a['equipment'] for a in ADDITIONS),
                         {'Sandbag': 16, 'Kółko do brzucha': 3, 'Masa własna': 24, 'Krzesło': 4})
        self.assertEqual(self.meta['exercises']['cw335']['tracking'], 'time')
        self.assertEqual(self.meta['exercises']['cw352']['tracking'], 'rounds')
        self.assertTrue(self.meta['exercises']['cw334']['unilateral'])
        self.assertFalse(self.meta['exercises']['cw333']['unilateral'])


if __name__ == '__main__':
    unittest.main()
