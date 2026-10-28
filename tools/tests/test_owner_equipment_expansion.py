"""Owner equipment: append-only raw data, explicit corrections and replay."""
import copy
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
import xlsx_to_json as converter
import validate_catalog

PREFIX_SHA = 'f08515f6c776162eeb448dcdfcf56b10d37c504a52083e5a8117aecd5a7ce02d'
IDS = {'cw009', 'cw123', 'cw135', 'cw136', 'cw165', 'cw267'}
PATHS = ['tools/source_data/baza_cwiczen_316.xlsx',
         'tools/source_data/exercise_names_pl.json',
         'tools/source_data/exercise_execution_conditions.json',
         'tools/source_data/catalog_sources.json', 'assets/data/exercises.json',
         'assets/data/exercise_catalog_metadata.json',
         'assets/data/exercise_catalog_corrections.json',
         'backend/app/mentor/exercise_catalog.json']


class OwnerEquipmentTests(unittest.TestCase):
    def test_reviewed_sources_distinguish_feet_hands_calf_and_eccentric_descent(self):
        from expand_catalog import ledger
        sources = ledger()
        urls = lambda keys: [sources['sources'][key]['url'] for key in keys]
        self.assertEqual(urls(sources['corrections']['cw009']), [
            'https://www.acefitness.org/resources/everyone/blog/5391/outdoor-functional-training-exercises/',
            'https://www.acefitness.org/about-ace/press-room/in-the-news/8284/what-muscles-do-pushups-work-tips-how-to-do-them-correctly-healthnews/'])
        self.assertEqual(urls(sources['corrections']['cw165']), [
            'https://exrx.net/WeightExercises/Gastrocnemius/BWStandingCalfRaise'])
        stepdown = sources['sources'][sources['exercises']['cw414']['source']]
        self.assertEqual(stepdown['url'],
                         'https://www.acefitness.org/resources/pros/expert-articles/5901/effective-exercises-to-target-the-glutes/')
        self.assertIn('Eccentric Step Down With Body Weight Only', stepdown['verifiedSections'])
        self.assertIn('reset', stepdown['verifiedSections'])
        self.assertIn('10–15 cm', stepdown['editorialAdaptation'])
        self.assertIn('10–12 inches', stepdown['editorialAdaptation'])
        self.assertIn('not measured', stepdown['editorialAdaptation'])
        feet = sources['sources'][sources['corrections']['cw009'][0]]
        self.assertIn('feet on bench', feet['verifiedSections'])
        self.assertIn('hands on ground', feet['verifiedSections'])
        calf = sources['sources'][sources['corrections']['cw165'][0]]
        self.assertIn('balance', calf['verifiedSections'])
        self.assertIn('optional dumbbell', calf['verifiedSections'])
        hand = sources['sources'][sources['exercises']['cw415']['source']]
        self.assertIn('Elevated-hand', hand['verifiedSections'])
        self.assertNotEqual(hand['url'], feet['url'])

    def test_stepdown_transfers_weight_to_floor_then_resets_not_support_leg_return(self):
        data = json.loads(converter.OUTPUT_PATH.read_text())
        row = data[413]
        self.assertEqual(row['zrodlo'],
                         'https://www.acefitness.org/resources/pros/expert-articles/5901/effective-exercises-to-target-the-glutes/')
        cues = ' '.join(row['kluczoweWskazowki'])
        self.assertIn('przenieś ciężar na podłogę', cues)
        self.assertIn('wejdź ponownie', cues)
        self.assertNotIn('Wróć pracą nogi podporowej', cues)
        self.assertNotIn('Odbijanie nogą dotykającą podłoża', row['czesteBledy'])

    def test_old_unreleased_refresh_cannot_modify_now_frozen_raw424(self):
        with tempfile.TemporaryDirectory(dir='/tmp/opencode') as tmp:
            root = Path(tmp)
            shutil.copytree(ROOT / 'tools', root / 'tools', ignore=shutil.ignore_patterns('__pycache__'))
            shutil.copytree(ROOT / 'assets/data', root / 'assets/data')
            path = root / PATHS[0]
            wb = openpyxl.load_workbook(path)
            ws = wb[converter.SHEET_NAME]
            ws.cell(415, 12).value = 'Old unreleased unsupported support-leg return'
            ws.cell(415, 13).value = 'Old unreleased error'
            ws.cell(415, 16).value = 'https://www.acefitness.org/resources/pros/expert-articles/6369/back-to-the-basics-how-to-create-innovative-fitness-programs/'
            wb.save(path)
            from update_exercise_names_pl import workbook_snapshot
            before = workbook_snapshot(wb)
            wb.close()
            command = [sys.executable, str(root / 'tools/expand_owner_equipment_catalog.py')]
            self.assertNotEqual(subprocess.run(command + ['--check'], capture_output=True).returncode, 0)
            self.assertNotEqual(subprocess.run(command, capture_output=True).returncode, 0)
            previous = path.read_bytes()
            result = subprocess.run(command + ['--refresh-unreleased-cw414'], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('cw001–cw424', result.stderr)
            wb = openpyxl.load_workbook(path)
            after = workbook_snapshot(wb)
            wb.close()
            changed = {(sheet, cell) for sheet, cells in before.items() for cell, value in cells.items()
                       if after[sheet].get(cell) != value}
            self.assertEqual(changed, set())
            self.assertEqual(path.read_bytes(), previous)

    def test_startup_corrections_are_exact_core_asset(self):
        # Source-level gate; the parent owns the Dart generatePackage fixture test.
        text = (ROOT / 'tools/generate_offline_package.dart').read_text()
        core = text.split('bool _core(String path) =>', 1)[1].split('].contains(path)', 1)[0]
        self.assertIn("'assets/assets/data/exercise_catalog_corrections.json'", core)
        self.assertNotIn("path.startsWith('assets/assets/data/')", text)

    def test_only_original_412_illustrations_remain_byte_identical(self):
        tree = subprocess.check_output(['git', 'ls-tree', '-r', '8f4ff6a', 'assets/exercises'], cwd=ROOT).decode()
        images = [line.split('\t') for line in tree.splitlines() if line.endswith('.webp')]
        self.assertEqual(len(images), 412)
        self.assertEqual(set(p.name for p in (ROOT / 'assets/exercises').glob('cw*.webp')
                             if int(p.stem[2:]) <= 412),
                         {f'cw{n:03d}.webp' for n in range(1, 413)})
        for info, path in images:
            content = (ROOT / path).read_bytes()
            blob = b'blob ' + str(len(content)).encode() + b'\0' + content
            self.assertEqual(hashlib.sha1(blob).hexdigest(), info.split()[2], path)

    def test_raw_prefix_and_exact_addition(self):
        data = json.loads(converter.OUTPUT_PATH.read_text())[:424]  # Owner publication scope.
        self.assertEqual(len(data), 424)
        self.assertEqual(hashlib.sha256(json.dumps(data[:412], ensure_ascii=False,
                                                  sort_keys=True).encode()).hexdigest(), PREFIX_SHA)
        self.assertEqual([e['id'] for e in data[412:424]], [f'cw{n}' for n in range(413, 425)])
        gears = [['Step']] * 3 + [['Ławka do brzuszków']] * 2 + [
            ['Ławka ujemna do wyciskania', 'Hantle']] + [['Poduszka równoważna']] * 2 + [
            ['Roller']] * 3 + [['Sandbag']]
        self.assertEqual([e['sprzet'] for e in data[412:424]], gears)
        self.assertIn('no_jumps', json.loads((ROOT / PATHS[5]).read_text())['exercises']['cw413']['conditions'])
        self.assertEqual(json.loads((ROOT / PATHS[5]).read_text())['exercises']['cw415']['positions'], ['floor'])
        self.assertNotIn('20 kg', data[-1]['progresja'])
        self.assertIn('nośność', ' '.join(data[417]['kluczoweWskazowki']))
        self.assertIn('niezweryfikowaną', ' '.join(data[417]['kluczoweWskazowki']))

    def test_correction_projection_is_explicit_and_does_not_mutate_raw(self):
        path = ROOT / PATHS[6]
        self.assertTrue(path.exists(), 'Missing versioned six-ID corrections')
        from effective_catalog import effective_catalog, validate_corrections
        raw = json.loads(converter.OUTPUT_PATH.read_text())
        before = copy.deepcopy(raw)
        manifest = json.loads(path.read_text())
        self.assertEqual(set(c['exerciseId'] for c in manifest['corrections']), IDS)
        validate_corrections(manifest)
        effective = effective_catalog(raw, manifest)
        self.assertEqual(raw, before)
        self.assertEqual(effective[122]['sprzet'], ['Wrist roller'])
        self.assertEqual(effective_catalog(effective, manifest), effective)
        for e, old in zip(effective, raw):
            if e['id'] not in IDS:
                self.assertEqual(e, old)
        for mutate in [lambda m: m['corrections'].append(m['corrections'][0]),
                       lambda m: m['corrections'][0]['fields'].update(ulubione=True),
                       lambda m: m['corrections'][0].update(revision=True),
                       lambda m: m['corrections'][0]['fields'].update(sprzet=[]),
                       lambda m: m['corrections'].pop()]:
            bad = copy.deepcopy(manifest)
            mutate(bad)
            with self.assertRaises(ValueError):
                validate_corrections(bad)

    def test_equipment_sets_are_complete_and_strict(self):
        meta = json.loads((ROOT / PATHS[5]).read_text())
        self.assertIn('equipmentSets', meta['exercises']['cw009'])
        self.assertEqual(meta['exercises']['cw009']['equipmentSets'], [
            ['Masa własna', 'Step'], ['Masa własna', 'Krzesło'], ['Masa własna', 'Ławeczka']])
        for eid in ['cw135', 'cw136', 'cw165']:
            self.assertEqual(meta['exercises'][eid]['equipmentSets'], [
                ['Masa własna', 'Step'], ['Hantle', 'Step'], ['Hantle', 'Ławeczka']])
        self.assertEqual(meta['exercises']['cw267']['equipmentSets'], [['Ławeczka'], ['Roller']])
        data = json.loads(converter.OUTPUT_PATH.read_text())
        for value in [[], [[]], ['Step'], [['Step', '']], [['Step', 1]], [['unknown']],
                      [['Step', 'Step']], [['Step'], ['Step']], None, 'Step']:
            bad = copy.deepcopy(meta)
            bad['exercises']['cw009']['equipmentSets'] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                validate_catalog.validate_metadata(data, bad)

    def test_mentor_projection_uses_corrections_and_sets(self):
        entries = {e['id']: e for e in json.loads((ROOT / PATHS[7]).read_text())['exercises']}
        self.assertEqual(entries['cw123']['equipment'], ['Wrist roller'])
        self.assertEqual(entries['cw009']['equipmentSets'][0], ['Masa własna', 'Step'])
        self.assertEqual(entries['cw418']['equipment'], ['Ławka ujemna do wyciskania', 'Hantle'])
        subprocess.run([sys.executable, str(ROOT / 'backend/tools/generate_mentor_catalog.py'), '--check'], check=True)

    def test_mentor_generator_remains_standard_library_only(self):
        # Backend's hash-locked Python environment does not install openpyxl.
        result = subprocess.run([sys.executable, '-S', str(ROOT / 'backend/tools/generate_mentor_catalog.py'),
                                 '--check'], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_replay_from_412_and_byte_idempotence(self):
        self.assertTrue((ROOT / 'tools/expand_owner_equipment_catalog.py').exists(), 'Missing owner append tool')
        with tempfile.TemporaryDirectory(dir='/tmp/opencode') as tmp:
            root = Path(tmp)
            shutil.copytree(ROOT / 'tools', root / 'tools', ignore=shutil.ignore_patterns('__pycache__'))
            (root / 'assets/data').mkdir(parents=True)
            (root / 'backend/tools').mkdir(parents=True)
            (root / 'backend/app/mentor').mkdir(parents=True)
            shutil.copy(ROOT / 'backend/tools/generate_mentor_catalog.py', root / 'backend/tools')
            for path in PATHS[:6]:
                (root / path).write_bytes(subprocess.check_output(['git', 'show', '8f4ff6a:' + path], cwd=ROOT))
            def run(tool, *args):
                subprocess.run([sys.executable, str(root / tool), *args], check=True, capture_output=True)
            tools = ['tools/expand_owner_equipment_catalog.py', 'tools/expand_bodyweight_catalog.py', 'tools/xlsx_to_json.py',
                     'tools/build_catalog_metadata.py', 'backend/tools/generate_mentor_catalog.py']
            for tool in tools:
                run(tool)
            for path in PATHS[1:]:
                self.assertEqual((root / path).read_bytes(), (ROOT / path).read_bytes(), path)
            hashes = [hashlib.sha256((root / p).read_bytes()).hexdigest() for p in PATHS]
            for tool in tools:
                run(tool)
            self.assertEqual(hashes, [hashlib.sha256((root / p).read_bytes()).hexdigest() for p in PATHS])
            run(tools[0], '--check')

    def test_converter_rejects_edit_inside_412_prefix(self):
        with tempfile.TemporaryDirectory(dir='/tmp/opencode') as tmp:
            path = Path(tmp) / 'edited.xlsx'
            wb = openpyxl.load_workbook(converter.XLSX_PATH)
            wb[converter.SHEET_NAME].cell(413, 3).value = 'Changed historical English name'
            wb.save(path)
            wb.close()
            with patch.object(converter, 'XLSX_PATH', path):
                with self.assertRaisesRegex(ValueError, 'cw001–cw412'):
                    converter.convert()


if __name__ == '__main__':
    unittest.main()
