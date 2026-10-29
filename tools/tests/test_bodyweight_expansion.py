"""Seven basic movements: frozen raw424, source coverage and append-only replay."""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools'))
import xlsx_to_json as converter
from update_exercise_names_pl import workbook_snapshot

BASE = 'acd52d8'
PREFIX_SHA = 'd734c9eb98c827f47d5755705fbd8819207fbe7431db65f26e61ec97c5e7c278'
PATHS = ['tools/source_data/baza_cwiczen_316.xlsx',
         'tools/source_data/exercise_names_pl.json',
         'tools/source_data/exercise_execution_conditions.json',
         'tools/source_data/catalog_sources.json', 'assets/data/exercises.json',
         'assets/data/exercise_catalog_metadata.json',
         'assets/data/exercise_catalog_corrections.json',
         'backend/app/mentor/exercise_catalog.json']
TOOLS = ['tools/expand_bodyweight_catalog.py', 'tools/xlsx_to_json.py',
         'tools/build_catalog_metadata.py', 'backend/tools/generate_mentor_catalog.py']
EN_PATH = 'tools/source_data/bodyweight_instructions_en.json'


class BodyweightTests(unittest.TestCase):
    def test_full_english_editorial_instructions_cover_exactly_seven_movements(self):
        path = ROOT / EN_PATH
        self.assertTrue(path.exists(), 'Missing complete English editorial sidecar')
        sidecar = json.loads(path.read_text())
        self.assertEqual(set(sidecar), {'schemaVersion', 'language', 'exercises'})
        self.assertEqual(sidecar['schemaVersion'], 1)
        self.assertEqual(sidecar['language'], 'en')
        self.assertEqual(set(sidecar['exercises']), {f'cw{n}' for n in range(425, 432)})
        raw = json.loads((ROOT / PATHS[4]).read_text())
        for pl in raw[424:431]:
            eid = pl['id']
            en = sidecar['exercises'][eid]
            self.assertEqual(set(en), {'name', 'description', 'cues', 'errors', 'tempo',
                                       'prescription', 'progression', 'regression'})
            self.assertEqual(en['name'], pl['nazwaEn'])
            for field in ('description', 'tempo', 'prescription', 'progression', 'regression'):
                self.assertIsInstance(en[field], str)
                self.assertTrue(en[field].strip(), (eid, field))
            self.assertGreaterEqual(len(en['description'].split()), 12, eid)
            for field, pl_field in [('cues', 'kluczoweWskazowki'), ('errors', 'czesteBledy')]:
                self.assertIsInstance(en[field], list)
                self.assertEqual(len(en[field]), len(pl[pl_field]), (eid, field))
                self.assertTrue(all(isinstance(s, str) and s.strip() for s in en[field]))
                self.assertTrue(all(s not in pl[pl_field] for s in en[field]), (eid, field))
            self.assertEqual(en['prescription'], '1-2 sets of 20-30 seconds' if eid in ('cw428', 'cw429')
                             else '1-2 sets of 5-8 repetitions per side' if eid in ('cw427', 'cw430', 'cw431')
                             else '1-2 sets of 5-8 repetitions')
            self.assertIn('Stop if you feel pain, numbness or dizziness.', en['cues'])

    def test_english_sidecar_is_generated_checked_and_idempotent_without_changing_catalog(self):
        with tempfile.TemporaryDirectory(dir='/tmp/opencode') as tmp:
            root = Path(tmp)
            shutil.copytree(ROOT / 'tools', root / 'tools', ignore=shutil.ignore_patterns('__pycache__'))
            # A completed431 device/repository, not a partially replayed workbook.
            for path in PATHS:
                (root / path).parent.mkdir(parents=True, exist_ok=True)
                (root / path).write_bytes(subprocess.check_output(['git', 'show', 'd778e69:' + path], cwd=ROOT))
            en_path = root / EN_PATH
            en_path.unlink(missing_ok=True)
            protected = {path: (root / path).read_bytes() for path in PATHS}
            command = [sys.executable, str(root / TOOLS[0])]
            check = subprocess.run(command + ['--check'], capture_output=True, text=True)
            self.assertNotEqual(check.returncode, 0, 'Missing EN must fail --check')
            self.assertFalse(en_path.exists(), '--check must not create the sidecar')
            for _ in range(2):
                result = subprocess.run(command, capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                self.assertTrue(en_path.exists(), 'Appender must generate full EN sidecar')
                self.assertEqual(en_path.read_bytes(), (ROOT / EN_PATH).read_bytes())
                self.assertEqual(subprocess.run(command + ['--check'], capture_output=True).returncode, 0)
                for path, before in protected.items():
                    self.assertEqual((root / path).read_bytes(), before, path)
            good = en_path.read_bytes()
            bad = json.loads(good)
            del bad['exercises']['cw431']['regression']
            en_path.write_text(json.dumps(bad))
            stale = en_path.read_bytes()
            self.assertNotEqual(subprocess.run(command + ['--check'], capture_output=True).returncode, 0)
            self.assertEqual(en_path.read_bytes(), stale, '--check must not repair stale EN')
            subprocess.run(command, check=True, capture_output=True)
            self.assertEqual(en_path.read_bytes(), good)
            for path, before in protected.items():
                self.assertEqual((root / path).read_bytes(), before, path)

    def test_exact_data_metadata_and_verified_source_coverage(self):
        data = json.loads((ROOT / PATHS[4]).read_text())
        self.assertEqual(len(data), 431)
        self.assertEqual([e['id'] for e in data], [f'cw{n:03d}' for n in range(1, 432)])
        self.assertEqual(hashlib.sha256(json.dumps(data[:424], ensure_ascii=False,
                                                  sort_keys=True).encode()).hexdigest(), PREFIX_SHA)
        self.assertEqual([e['nazwaPl'] for e in data[424:]], [
            'Pompki od ściany', 'Przysiad bez obciążenia', 'Zakroki bez obciążenia',
            'Marsz w miejscu', 'Marsz w siadzie', 'Prostowanie kolana w siadzie bez gumy',
            'Odwodzenie ugiętej nogi w bok w klęku podpartym bez gumy'])
        self.assertEqual([e['nazwaEn'] for e in data[424:]], [
            'Wall push-up', 'Bodyweight squat', 'Bodyweight reverse lunge',
            'March in place', 'Seated march', 'Seated knee extension', 'Bodyweight fire hydrant'])
        self.assertEqual([e['sprzet'] for e in data[424:]],
                         [['Masa własna']] * 4 + [['Krzesło']] * 2 + [['Masa własna']])
        meta = json.loads((ROOT / PATHS[5]).read_text())
        self.assertEqual(set(meta['exercises']), {e['id'] for e in data})
        ledger = json.loads((ROOT / PATHS[3]).read_text())
        mentor = {e['id']: e for e in json.loads((ROOT / PATHS[7]).read_text())['exercises']}
        old_meta = json.loads(subprocess.check_output(['git', 'show', BASE + ':' + PATHS[5]], cwd=ROOT))
        for eid, old in old_meta['exercises'].items():
            self.assertEqual(meta['exercises'][eid], old)
        for i, e in enumerate(data[424:]):
            m = meta['exercises'][e['id']]
            self.assertEqual(m['tracking'], 'time' if i in (3, 4) else 'reps')
            self.assertEqual(m['unilateral'], i in (2, 5, 6))
            self.assertEqual(m['positions'], ['standing' if i < 4 else 'seated' if i < 6 else 'floor'])
            expected = {'no_jumps', 'quiet', 'small_space'}
            if i != 6:
                expected |= {'no_kneeling', 'no_floor'}
            if i in (1, 2, 3, 5):
                expected.add('no_hand_support')
            self.assertEqual(set(m['conditions']), expected)
            source = ledger['sources'][ledger['exercises'][e['id']]['source']]
            self.assertEqual(source['url'], e['zrodlo'])
            self.assertTrue(source['url'].startswith('https://'))
            self.assertTrue(source['verifiedSections'])
            self.assertEqual(source['readOn'], '2026-10-28')
            self.assertEqual(mentor[e['id']]['equipment'], e['sprzet'])
            for k in ('kluczoweWskazowki', 'czesteBledy', 'tempo', 'seriexPowtorzenia', 'progresja', 'regresja'):
                self.assertTrue(e[k], (e['id'], k))

    def test_all_424_images_match_pinned_baseline(self):
        tree = subprocess.check_output(['git', 'ls-tree', '-r', BASE, 'assets/exercises'], cwd=ROOT).decode()
        images = [line.split('\t') for line in tree.splitlines() if line.endswith('.webp')]
        self.assertEqual(len(images), 424)
        self.assertEqual({path for _, path in images},
                         {f'assets/exercises/cw{n:03d}.webp' for n in range(1, 425)})
        self.assertEqual({p.name for p in (ROOT / 'assets/exercises').glob('*.webp')},
                         {f'cw{n:03d}.webp' for n in range(1, 432)})
        for info, path in images:
            content = (ROOT / path).read_bytes()
            self.assertEqual(hashlib.sha1(b'blob ' + str(len(content)).encode() + b'\0' + content).hexdigest(),
                             info.split()[2], path)

    def test_converter_guards_all_raw424(self):
        with tempfile.TemporaryDirectory(dir='/tmp/opencode') as tmp:
            path = Path(tmp) / 'edited.xlsx'
            wb = openpyxl.load_workbook(converter.XLSX_PATH)
            wb[converter.SHEET_NAME].cell(425, 3).value = 'Changed protected carry'
            wb.save(path)
            wb.close()
            with patch.object(converter, 'XLSX_PATH', path):
                with self.assertRaisesRegex(ValueError, 'cw001–cw424'):
                    converter.convert()

    def test_replay_preserves_cells_and_all_generated_bytes_on_second_run(self):
        self.assertTrue((ROOT / TOOLS[0]).exists(), 'Missing bodyweight append CLI')
        with tempfile.TemporaryDirectory(dir='/tmp/opencode') as tmp:
            root = Path(tmp)
            shutil.copytree(ROOT / 'tools', root / 'tools', ignore=shutil.ignore_patterns('__pycache__'))
            for path in PATHS:
                (root / path).parent.mkdir(parents=True, exist_ok=True)
                (root / path).write_bytes(subprocess.check_output(['git', 'show', BASE + ':' + path], cwd=ROOT))
            (root / 'backend/tools').mkdir(parents=True)
            shutil.copy(ROOT / TOOLS[-1], root / TOOLS[-1])
            wb = openpyxl.load_workbook(root / PATHS[0])
            before = workbook_snapshot(wb)
            wb.close()
            def run(tool, *args):
                result = subprocess.run([sys.executable, str(root / tool), *args], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            for tool in TOOLS:
                run(tool)
            wb = openpyxl.load_workbook(root / PATHS[0])
            after = workbook_snapshot(wb)
            wb.close()
            for sheet, cells in before.items():
                for coord, value in cells.items():
                    self.assertEqual(after[sheet].get(coord), value, f'{sheet}!{coord}')
            for path in PATHS[1:]:
                self.assertEqual((root / path).read_bytes(), (ROOT / path).read_bytes(), path)
            hashes = [hashlib.sha256((root / p).read_bytes()).hexdigest() for p in PATHS]
            for tool in TOOLS + ['tools/expand_catalog.py', 'tools/expand_ez_catalog.py',
                                  'tools/expand_owner_equipment_catalog.py']:
                run(tool)
            self.assertEqual(hashes, [hashlib.sha256((root / p).read_bytes()).hexdigest() for p in PATHS])
            for tool in [TOOLS[0], 'tools/expand_catalog.py', 'tools/expand_ez_catalog.py',
                         'tools/expand_owner_equipment_catalog.py', TOOLS[2], TOOLS[3]]:
                run(tool, '--check')
            self.assertEqual(hashes, [hashlib.sha256((root / p).read_bytes()).hexdigest() for p in PATHS])


if __name__ == '__main__':
    unittest.main()
