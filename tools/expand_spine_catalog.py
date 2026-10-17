#!/usr/bin/env python3
"""Append the second reviewed expansion; never edit exported exercises.json.

Idempotent; --check writes nothing. Existing cells in every sheet are protected.
Execution decisions are appended to the full static 363-entry input, not inferred.
"""
import argparse
from copy import copy
import json
from pathlib import Path
import tempfile

import openpyxl

import xlsx_to_json as converter
from build_catalog_metadata import MUSCLES
from build_execution_conditions import INPUT_PATH, validate_classification
from expand_catalog import SOURCE_PATH, ledger
from spine_additions import ADDITIONS, EXPECTED_COUNT, SOURCES
from update_exercise_names_pl import MAPPING_PATH, workbook_snapshot

LABELS = {i: label for i, label, _ in MUSCLES}


def workbook_row(a):
    return [a['number'], a['pl'], a['en'],
            ', '.join(LABELS[m] for m in a['primary']),
            ', '.join(LABELS[m] for m in a['secondary']) or '-',
            a['equipment'], a['kind'], a['level'], a['pattern'],
            a['prescription'], a['tempo'], a['cues'], a['errors'],
            a['progression'], a['regression'], SOURCES[a['source']]['url']]


def checked_append(mapping, key, value):
    if key in mapping and mapping[key] != value:
        raise ValueError(f'Existing value differs: {key}')
    mapping[key] = value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    before = converter.convert()
    if len(before) not in (363, EXPECTED_COUNT):
        raise ValueError('Expected baseline 363 or completed spine expansion')
    names = json.loads(MAPPING_PATH.read_text(encoding='utf-8'))
    execution = json.loads(INPUT_PATH.read_text(encoding='utf-8'))
    validate_classification(execution, {e['id'] for e in before})
    wb = openpyxl.load_workbook(converter.XLSX_PATH)
    original = workbook_snapshot(wb)
    ws = wb[converter.SHEET_NAME]
    existing = {row[0]: list(row) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    added = 0
    for a in ADDITIONS:
        n = a['number']
        row = workbook_row(a)
        if n in existing:
            if existing[n] != row:
                raise ValueError(f'Existing expansion row differs: cw{n}')
        elif args.check:
            raise ValueError(f'Missing spine row cw{n}')
        else:
            ws.append(row)
            added += 1
            for cell, template in zip(ws[ws.max_row], ws[2]):
                cell._style = copy(template._style)
        eid = f'cw{n}'
        if args.check and (eid not in names or eid not in execution['exercises']):
            raise ValueError(f'Missing static input: {eid}')
        checked_append(names, eid, a['pl'])
        checked_append(execution['exercises'], eid,
                       {k: a[k] for k in ('conditions', 'positions')})
    validate_classification(execution, {f'cw{n:03d}' for n in range(1, EXPECTED_COUNT + 1)})
    outputs = {MAPPING_PATH: names, INPUT_PATH: execution, SOURCE_PATH: ledger()}
    texts = {path: json.dumps(value, ensure_ascii=False, indent=2) + '\n'
             for path, value in outputs.items()}
    if args.check:
        for path, text in texts.items():
            if path.read_text(encoding='utf-8') != text:
                raise ValueError(f'Spine input differs: {path}')
        print(f'Spine CHECK OK: {len(ADDITIONS)} rows, names, explicit conditions, sources')
    else:
        # Avoid rewriting the ZIP workbook (timestamps) on a no-op second run.
        if added:
            with tempfile.NamedTemporaryFile(dir=converter.XLSX_PATH.parent,
                                             suffix='.xlsx', delete=False) as tmp:
                path = Path(tmp.name)
            try:
                wb.save(path)
                saved = openpyxl.load_workbook(path)
                snapshot = workbook_snapshot(saved)
                saved.close()
                for sheet, cells in original.items():
                    for coord, value in cells.items():
                        if snapshot[sheet].get(coord) != value:
                            raise ValueError(f'Existing cell changed: {sheet}!{coord}')
                original_path = converter.XLSX_PATH
                try:
                    converter.XLSX_PATH = path
                    after = converter.convert()
                finally:
                    converter.XLSX_PATH = original_path
                if after[:len(before)] != before or len(after) != EXPECTED_COUNT:
                    raise ValueError('Existing records changed or unexpected new count')
                path.chmod(original_path.stat().st_mode & 0o777)
                path.replace(original_path)
            finally:
                path.unlink(missing_ok=True)
        for path, text in texts.items():
            if path.read_text(encoding='utf-8') != text:
                path.write_text(text, encoding='utf-8')
        print(f'Spine: added {added} workbook rows; run xlsx_to_json.py and build_catalog_metadata.py')
    wb.close()


if __name__ == '__main__':
    main()
