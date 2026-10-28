#!/usr/bin/env python3
"""Append reviewed rows to BAZA_GLOWNA, preserving every existing cell.

Idempotent. Does not write exercises.json: run xlsx_to_json.py afterwards.
The historical workbook filename is retained to avoid breaking existing tools.
"""
import argparse
from copy import copy
import json
from pathlib import Path
import tempfile

import openpyxl

from build_catalog_metadata import MUSCLES
from catalog_additions import ADDITIONS, SOURCES
from spine_additions import SOURCES as SPINE_SOURCES, source_exercises
import xlsx_to_json as converter
from update_exercise_names_pl import MAPPING_PATH, workbook_snapshot

SOURCE_PATH = Path(__file__).parent / 'source_data/catalog_sources.json'
LABELS = {i: label for i, label, _ in MUSCLES}


def workbook_row(a):
    side = ' na stronę' if a['unilateral'] else ''
    prescription = ('2 x 20-30 sek' if a['tracking'] == 'time' else
                    '3 rundy x 30 sek pracy' if a['tracking'] == 'rounds' else
                    '2 x 8-10' if 'tai_chi' in a['style'] else '3 x 6-10') + side
    return [a['number'], a['pl'], a['en'],
            ', '.join(LABELS[m] for m in a['primary']),
            ', '.join(LABELS[m] for m in a['secondary']) or '-',
            a['equipment'], a['kind'], a['level'],
            'Flow' if 'tai_chi' in a['style'] else 'Izometria' if a['tracking'] == 'time' else 'Ruch złożony',
            prescription, 'Wolno i kontrolowanie, fazy dynamiczne bez szarpania',
            a['cues'], a['errors'], a['progression'], a['regression'], SOURCES[a['source']][0]]


def ledger():
    from partner_additions import SOURCES as PARTNER_SOURCES, source_records
    from ez_additions import SOURCES as EZ_SOURCES, VERIFIED, source_records as ez_records
    result = {
        'reviewedOn': '2026-10-16',
        'method': 'Read publisher text; original Polish descriptions; conservative equipment adaptation, no copied prose. Source validates movement, not individualized load or medical suitability.',
        'equipment': {
            'Sandbag': {
                'url': 'https://bushido-sport.pl/product-pol-11-SANDBAG-BUSHIDO-SAND-BAG-CROSSTRAINING-FITNESS-35-KG.html',
                'verified': '7 handles, 3 removable filler bags, maximum recommended 35 kg; user-owned 20 kg is context only, never a prescription. DVRT bag shapes differ: use a lighter fill and secure handles; no throws.'},
            'Kółko do brzucha': {
                'url': 'https://www.sport-shop.pl/walek-kolko-podwojne-do-cwiczen-mata-pod-kolana-gymtek-podwojne-mata-czarne-p-170878.html',
                'verified': 'Double wheel, two side handles, knee mat; no foot straps or assisted return listed. Both hands stay on handles; no feet-on-wheel or single-arm exercises.'}},
        'contextOnly': [{
            'url': 'https://news.va.gov/107780/live-whole-health-134-seated-tai-chi-moves/',
            'verified': '18-minute seated Tai Chi session by Lori Enloe; article supports accessibility context, not a textual movement-by-movement attribution.'}],
        'sources': {key: {'url': url, 'verifiedSections': sections}
                    for key, (url, sections) in SOURCES.items()},
        'exercises': {f'cw{a["number"]:03d}': {'source': a['source'], 'movement': a['en']}
                      for a in ADDITIONS},
        'excluded': [
            'Ab wheel: partial/wall-limited rollout is regression, standing rollout is advanced progression, not extra IDs. Foot straps/auto-return/single central handle are unavailable.',
            'Yoga bridge duplicates cw154; locust duplicates cw185. Existing bear crawl cw223, crab walk cw224, L-sit cw203, dead bug cw179, inchworm cw247 are not re-added.',
            'No speculative Animal Rx variants inferred from a list of names; 8 documented calisthenics movements instead of 15.',
            'Seated opening/clouds already cw282–289; preserve cw283/cw285 despite similarity. No extra opening/cloud variants.',
            'No karate.'
        ],
    }
    result['sources'].update(SPINE_SOURCES)
    result['exercises'].update(source_exercises())
    result['sources'].update({
        'partner_expansion_' + key: {
            'url': url,
            'verifiedSections': '2026-10-20 conversation research: named movement instructions/snippets; original Polish adaptation, not medical suitability or individualized dosage.',
        } for key, url in PARTNER_SOURCES.items()
    })
    result['exercises'].update(source_records())
    result['sources'].update({
        'ez_expansion_' + key: {'url': url, 'verifiedSections': VERIFIED[key]}
        for key, url in EZ_SOURCES.items()
    })
    result['exercises'].update(ez_records())
    result['equipment'].update({
        'Gryf łamany EZ': {
            'url': EZ_SOURCES['ez_curl'],
            'verified': 'Angled EZ bar grip; load logged as total bar plus plates, not per hand. Secure collars; equipment limit is not prescribed load.'},
        'Rower stacjonarny': {
            'url': EZ_SOURCES['stationary_bike'],
            'verified': 'Stationary bike with adjusted seat and handlebar reach; slight knee bend, easy seated pedaling. Not a treadmill or outdoor-bike substitute.'},
    })
    from owner_equipment_additions import SOURCES as OWNER_SOURCES, VERIFIED as OWNER_VERIFIED, EDITORIAL_ADAPTATIONS, source_records as owner_records
    result['sources'].update({
        'owner_equipment_' + key: {'url': url, 'verifiedSections': '2026-10-22 supplied publisher research: ' + OWNER_VERIFIED[key]}
        for key, url in OWNER_SOURCES.items()
    })
    result['exercises'].update(owner_records())
    for key, note in EDITORIAL_ADAPTATIONS.items():
        result['sources']['owner_equipment_' + key]['editorialAdaptation'] = note
    result['corrections'] = {
        'cw009': ['owner_equipment_feet_pushup', 'owner_equipment_feet_pushup_low'],
        'cw123': ['owner_equipment_wrist'],
        'cw135': ['owner_equipment_step_up', 'owner_equipment_step_glutes'],
        'cw136': ['owner_equipment_step_glutes'],
        'cw165': ['owner_equipment_standing_calf'],
        'cw267': ['owner_equipment_thoracic'],
    }
    result['ownerEquipmentContext'] = {
        'Step': 'Adjustable base; height not measured, 10–15 cm only illustrative, declared 150 kg not independently verified. No jumps or loaded sandbag step-ups.',
        'Ławka do brzuszków': 'Decline and leg lock confirmed; suitability/load rating for dumbbell pressing unverified. Not an alias for generic Ławeczka or Ławka ujemna do wyciskania.',
        'Sandbag': 'Adjustable but owner operationally keeps 20 kg. Context only, never a universal starting load or progression target.',
        'Pasy do zwisu za nogi': 'Inversion leg straps, not TRX or arm slings; no added inversion movements.',
        'Wrist roller': 'Handle, cord and suspended plate; no pulley station and no dumbbell pair.',
    }
    from bodyweight_additions import SOURCES as BW_SOURCES, VERIFIED as BW_VERIFIED, ADAPTATIONS, source_records as bw_records
    result['sources'].update({
        'bodyweight_' + key: {'url': url, 'verifiedSections': BW_VERIFIED[key],
                             'readOn': '2026-10-28', 'editorialAdaptation': ADAPTATIONS[key]}
        for key, url in BW_SOURCES.items()
    })
    result['exercises'].update(bw_records())
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    before = converter.convert()
    wb = openpyxl.load_workbook(converter.XLSX_PATH)
    ws = wb[converter.SHEET_NAME]
    original = workbook_snapshot(wb)
    existing = {row[0]: list(row) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    names = json.loads(MAPPING_PATH.read_text())
    added = 0
    for a in ADDITIONS:
        row = workbook_row(a)
        if a['number'] in existing:
            if existing[a['number']] != row:
                raise ValueError(f'Existing expansion ID differs: {a["number"]}')
        elif args.check:
            raise ValueError(f'Missing expansion ID: {a["number"]}')
        else:
            ws.append(row)
            added += 1
            for cell, template in zip(ws[ws.max_row], ws[2]):
                cell._style = copy(template._style)
        eid = f'cw{a["number"]:03d}'
        if eid in names and names[eid] != a['pl']:
            raise ValueError(f'Existing mapping differs: {eid}')
        if args.check and eid not in names:
            raise ValueError(f'Missing mapped name: {eid}')
        names[eid] = a['pl']
    text = json.dumps(ledger(), ensure_ascii=False, indent=2) + '\n'
    if args.check:
        if SOURCE_PATH.read_text() != text:
            raise ValueError('Source ledger differs from reviewed input')
        print(f'Expansion CHECK OK: {len(ADDITIONS)} rows, names and sources')
    elif added:
        with tempfile.NamedTemporaryFile(dir=converter.XLSX_PATH.parent, suffix='.xlsx', delete=False) as tmp:
            path = Path(tmp.name)
        try:
            wb.save(path)
            saved = openpyxl.load_workbook(path)
            snapshot = workbook_snapshot(saved)
            saved.close()
            for sheet, cells in original.items():
                for coord, value in cells.items():
                    if snapshot[sheet].get(coord) != value:
                        raise ValueError(f'Existing cell changed: {sheet}!{coord}')
            original_path = converter.XLSX_PATH
            try:
                converter.XLSX_PATH = path
                after = converter.convert()
            finally:
                converter.XLSX_PATH = original_path
            if after[:len(before)] != before:
                raise ValueError('Existing export records changed')
            path.replace(original_path)
        finally:
            path.unlink(missing_ok=True)
        MAPPING_PATH.write_text(json.dumps(names, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        SOURCE_PATH.write_text(text, encoding='utf-8')
        print(f'Workbook: {len(after)} rows; added {len(after) - len(before)}. Run xlsx_to_json.py.')
    else:
        # Do not re-zip an unchanged workbook: ZIP timestamps are not stable.
        for path, value in [(MAPPING_PATH, json.dumps(names, ensure_ascii=False, indent=2) + '\n'),
                            (SOURCE_PATH, text)]:
            if path.read_text() != value:
                path.write_text(value, encoding='utf-8')
        print(f'Expansion OK: {len(before)} rows; no workbook changes')
    wb.close()


if __name__ == '__main__':
    main()
